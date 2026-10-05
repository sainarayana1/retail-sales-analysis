"""Build outputs/Sales_Analysis.xlsx : cleaned Data sheet + Summary sheet using live SUMIFS formulas."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

df = pd.read_csv("data/superstore_clean.csv", parse_dates=["Order Date"])
keep = ["Order ID","Order Date","Ship Mode","Customer Name","Segment","Region","State","Category","Sub-Category",
        "Sales","Quantity","Discount","Profit","Year"]
df = df[keep]

wb = Workbook()
ws = wb.active; ws.title = "Data"
ws.append(keep)
for r in df.itertuples(index=False):
    ws.append([v.to_pydatetime() if hasattr(v, "to_pydatetime") else v for v in r])
n = len(df) + 1
hdr = PatternFill("solid", fgColor="1F4E78")
for c in ws[1]:
    c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = hdr
for i in range(1, len(keep) + 1):
    ws.column_dimensions[get_column_letter(i)].width = 16
for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    row[0].number_format = "yyyy-mm-dd"
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions

S = wb.create_sheet("Summary", 0)
f = lambda **k: Font(name="Arial", **k)
S["A1"] = "Retail Sales Performance - Summary"; S["A1"].font = f(bold=True, size=14)
S["A2"] = "All figures are live formulas over the 'Data' sheet (SUMIFS / COUNTA)."; S["A2"].font = f(italic=True, color="666666")
Rng = lambda col: f"Data!${col}$2:${col}${n}"

for c, h in zip("ABCD", ["Total Sales", "Total Profit", "Order lines", "Profit Margin"]):
    S[f"{c}4"] = h; S[f"{c}4"].font = f(bold=True, color="FFFFFF"); S[f"{c}4"].fill = hdr
S["A5"] = f"=SUM({Rng('J')})"; S["B5"] = f"=SUM({Rng('M')})"
S["C5"] = f"=COUNTA({Rng('A')})"; S["D5"] = "=B5/A5"
for c in "ABCD": S[f"{c}5"].font = f(bold=True, size=12)
for c in "AB": S[f"{c}5"].number_format = "$#,##0"
S["C5"].number_format = "#,##0"; S["D5"].number_format = "0.0%"

def table(top, title, dim, labels, keycol):
    S.cell(top, 1, title).font = f(bold=True, size=12)
    for j, h in enumerate([dim, "Sales", "Profit", "Margin"], 1):
        c = S.cell(top + 1, j, h); c.font = f(bold=True, color="FFFFFF"); c.fill = hdr
    for i, lab in enumerate(labels):
        r = top + 2 + i
        S.cell(r, 1, lab).font = f()
        S.cell(r, 2, f"=SUMIFS({Rng('J')},{Rng(keycol)},A{r})").number_format = "$#,##0"
        S.cell(r, 3, f"=SUMIFS({Rng('M')},{Rng(keycol)},A{r})").number_format = "$#,##0"
        S.cell(r, 4, f"=IF(B{r}=0,0,C{r}/B{r})").number_format = "0.0%"
        for c in (2, 3, 4): S.cell(r, c).font = f()
    return top + 2 + len(labels) + 1

nxt = table(7, "Sales by Category", "Category", sorted(df["Category"].unique()), "H")
nxt = table(nxt, "Sales by Region", "Region", sorted(df["Region"].unique()), "F")
nxt = table(nxt, "Sales by Segment", "Segment", sorted(df["Segment"].unique()), "E")
table(nxt, "Sales by Year", "Year", [int(y) for y in sorted(df["Year"].unique())], "N")
for col, w in zip("ABCD", (24, 16, 20, 14)): S.column_dimensions[col].width = w
wb.save("outputs/Sales_Analysis.xlsx"); print("Saved outputs/Sales_Analysis.xlsx")
